"""Shared document-building helpers: BibTeX parsing, numbered citations, Markdown -> DOCX, DOCX -> PDF.

Markdown subset understood (used by docs/report_src/report.md):
  # Title            ## Section        ### Subsection
  paragraphs with **bold**, *italic*, `code`, citations written as [@key] or [@key1; @key2]
  - bullets, 1. numbered items, pipe tables (optionally preceded by a `Table: caption` line)
  ![caption](relative/or/absolute/path.png)      figures (auto numbered)
  $$ latex $$                                     display equations (rendered with matplotlib mathtext, auto numbered)
  {{table:file.csv|caption|compact}}              measured table from artifacts/results/tables
  <!-- pagebreak -->
"""
from __future__ import annotations

import csv
import io
import re
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(0x0B, 0x25, 0x45)
TEAL = RGBColor(0x13, 0x80, 0x8B)
ROOT = Path(__file__).resolve().parents[1]


# ----------------------------------------------------------------------------------------- bibtex
_ACCENTS = {("'", "c"): "ć", ("'", "e"): "é", ("'", "o"): "ó", ("'", "E"): "É", ("`", "o"): "ò", ("v", "c"): "č",
            ("'", "a"): "á", ("'", "i"): "í", ("'", "u"): "ú", ('"', "o"): "ö", ('"', "u"): "ü", ('"', "a"): "ä"}


def _unlatex(v):
    """Turn LaTeX accent commands in BibTeX values into Unicode characters and unescape ampersands."""
    def rep(m):
        return _ACCENTS.get((m.group(1), m.group(2)), m.group(2))
    v = re.sub(r"\{?[\\]([\'`\"^~]|v)\s*\{?([A-Za-z])\}?\}?", rep, v)
    return v.replace("\\&", "&")


def parse_bib(path):
    text = Path(path).read_text(encoding="utf-8")
    text = "\n".join(l for l in text.splitlines() if not l.lstrip().startswith("%"))
    entries = {}
    for m in re.finditer(r"@(\w+)\s*\{\s*([^,\s]+)\s*,", text):
        etype, key = m.group(1).lower(), m.group(2)
        i = m.end()
        depth = 1
        j = i
        while j < len(text) and depth:
            depth += {"{": 1, "}": -1}.get(text[j], 0)
            j += 1
        body = text[i:j - 1]
        fields = {}
        for fm in re.finditer(r"(\w+)\s*=\s*", body):
            k = fm.group(1).lower()
            s = fm.end()
            if s >= len(body):
                continue
            if body[s] == "{":
                d, e = 1, s + 1
                while e < len(body) and d:
                    d += {"{": 1, "}": -1}.get(body[e], 0)
                    e += 1
                v = body[s + 1:e - 1]
            elif body[s] == '"':
                e = body.index('"', s + 1)
                v = body[s + 1:e]
            else:
                e = s
                while e < len(body) and body[e] not in ",\n":
                    e += 1
                v = body[s:e]
            fields.setdefault(k, re.sub(r"\s+", " ", _unlatex(v).replace("{", "").replace("}", "")).strip())
        entries[key] = {"type": etype, **fields}
    return entries


def format_ref(e):
    authors = e.get("author", "")
    names = [a.strip() for a in authors.split(" and ")] if authors else []

    def short(n):
        if "," in n:
            last, first = [x.strip() for x in n.split(",", 1)]
            ini = " ".join(w[0] + "." for w in first.replace("-", " ").split() if w)
            return f"{ini} {last}".strip()
        return n

    names = [short(n) for n in names]
    if len(names) > 4:
        au = ", ".join(names[:3]) + ", et al."
    elif len(names) > 1:
        au = ", ".join(names[:-1]) + " and " + names[-1]
    else:
        au = names[0] if names else ""
    title = e.get("title", "")
    venue = e.get("booktitle") or e.get("journal") or e.get("publisher") or e.get("howpublished") or e.get("institution") or ""
    if e.get("type") == "misc" and e.get("eprint"):
        venue = venue or f"arXiv:{e['eprint']}"
    parts = [f"{au}, \"{title},\"" if au else f"\"{title},\""]
    if venue:
        parts.append(venue + ",")
    if e.get("volume"):
        parts.append(f"vol. {e['volume']}" + (f", no. {e['number']}" if e.get("number") else "") + ",")
    if e.get("pages"):
        parts.append(f"pp. {e['pages'].replace('--', '-')},")
    parts.append(f"{e.get('year', '')}.")
    link = ""
    if e.get("doi"):
        link = f" doi:{e['doi']}"
    elif e.get("eprint"):
        link = f" arXiv:{e['eprint']}"
    elif e.get("url"):
        link = f" {e['url']}"
    return " ".join(parts) + link


class Citations:
    def __init__(self, bib):
        self.bib = bib
        self.order = []

    def num(self, key):
        if key not in self.bib:
            raise KeyError(f"citation key {key!r} not in references.bib")
        if key not in self.order:
            self.order.append(key)
        return self.order.index(key) + 1

    def sub(self, text):
        def rep(m):
            keys = [k.strip().lstrip("@") for k in m.group(1).split(";")]
            nums = sorted(self.num(k) for k in keys)
            return "[" + ", ".join(str(n) for n in nums) + "]"
        return re.sub(r"\[(@[^\]]+)\]", rep, text)

    def entries(self):
        return [(i + 1, k, format_ref(self.bib[k])) for i, k in enumerate(self.order)]


# ----------------------------------------------------------------------------------------- tables
def csv_to_rows(path, compact=True, drop=()):
    with open(path, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    if compact:
        rows = [[re.sub(r" \[n=\d+\]", "", c) for c in r] for r in rows]
    if drop:
        idx = [i for i, h in enumerate(rows[0]) if h not in drop]
        rows = [[r[i] for i in idx] for r in rows]
    return rows


def rows_to_md(rows):
    head = "| " + " | ".join(rows[0]) + " |"
    sep = "|" + "|".join("---" for _ in rows[0]) + "|"
    body = ["| " + " | ".join(c.replace("|", "/") for c in r) + " |" for r in rows[1:]]
    return "\n".join([head, sep] + body)


# ----------------------------------------------------------------------------------------- equations
def render_equation(latex, path, fs=15):
    fig = plt.figure(figsize=(0.1, 0.1))
    fig.text(0, 0, f"${latex}$", fontsize=fs, color="#0b2545")
    fig.savefig(path, dpi=220, bbox_inches="tight", pad_inches=0.05, transparent=True)
    plt.close(fig)


# ----------------------------------------------------------------------------------------- docx
def _shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_)
    tcPr.append(shd)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _add_field(par, instr):
    r = par.add_run()
    for t, txt in (("begin", None), (None, instr), ("end", None)):
        if t:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), t)
            r._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            r._r.append(it)


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")


def add_inline(par, text, size=None, color=None, italic=False):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt((size or 11) - 1)
            continue
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1])
            r.italic = True
        else:
            r = par.add_run(tok)
        if italic:
            r.italic = True
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color


def new_document(title_for_footer=""):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)  # A4
    sec.left_margin = sec.right_margin = Inches(0.95)
    sec.top_margin = sec.bottom_margin = Inches(0.9)
    st = doc.styles["Normal"]
    st.font.name = "Cambria"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria")
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.12
    for name, size, before in (("Heading 1", 15, 16), ("Heading 2", 12.5, 12), ("Heading 3", 11, 8)):
        h = doc.styles[name]
        h.font.name = "Calibri"
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = NAVY
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(5)
        h.paragraph_format.keep_with_next = True
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(f"{title_for_footer}   |   page ")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    _add_field(fp, "PAGE")
    return doc


def add_table(doc, rows, caption=None, font=7.5, col_widths=None, number=None):
    if caption:
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(8)
        add_inline(p, f"**Table {number}.** {caption}", size=9, color=NAVY)
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(rows):
        for j, c in enumerate(r):
            cell = t.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            add_inline(par, c, size=font, color=RGBColor(255, 255, 255) if i == 0 else None)
            if i == 0:
                for run in par.runs:
                    run.bold = True
                _shade(cell, "0B2545")
            elif i % 2 == 0:
                _shade(cell, "EEF3F7")
    _repeat_header(t.rows[0])
    if not col_widths:
        total = 6.35
        wts = []
        for j in range(len(rows[0])):
            longest = max(len(str(r[j])) for r in rows)
            wts.append(min(max(longest, 6), 44))
        col_widths = [max(0.45, total * w / sum(wts)) for w in wts]
        scale = total / sum(col_widths)
        col_widths = [w * scale for w in col_widths]
    t.autofit = False
    for row in t.rows:
        for j, w in enumerate(col_widths):
            row.cells[j].width = Inches(w)
    # keep the whole table on one page when it fits; never split a row
    for ri, row in enumerate(t.rows):
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        cs.set(qn("w:val"), "true")
        trPr.append(cs)
        if ri < len(t.rows) - 1:
            for cell in row.cells:
                for par in cell.paragraphs:
                    par.paragraph_format.keep_with_next = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def add_figure(doc, path, caption, number, width=6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(c, f"**Figure {number}.** {caption}", size=9, color=NAVY)


def add_equation(doc, img, number):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    c0, c1 = t.rows[0].cells
    c0.width, c1.width = Inches(5.6), Inches(0.6)
    p = c0.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(img), height=Inches(0.34) if Path(img).stat().st_size < 12000 else Inches(0.42))
    q = c1.paragraphs[0]
    q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    q.paragraph_format.space_after = Pt(0)
    q.add_run(f"({number})").font.size = Pt(10)
    for c in (c0, c1):
        c.vertical_alignment = 1


def docx_to_pdf(docx_path, pdf_path):
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        doc = word.Documents.Open(str(Path(docx_path).resolve()), ReadOnly=True)
        doc.Fields.Update()
        doc.SaveAs2(str(Path(pdf_path).resolve()), FileFormat=17)
        pages = doc.ComputeStatistics(2)
        doc.Close(False)
    finally:
        word.Quit()
    return pages
